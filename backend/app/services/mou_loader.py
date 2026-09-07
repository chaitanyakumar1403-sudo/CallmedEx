"""
CallMedex Authentic MOU Loader Service
=======================================
Extracts and provides full unabridged legal agreements directly from the 
canonical `mous/` document repository for all CallMedex healthcare provider roles.
"""
import os
import logging
from typing import Dict, Optional

logger = logging.getLogger(__name__)

# Search candidates for the mous directory
POSSIBLE_MOU_DIRS = [
    os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..", "mous")),
    os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "mous")),
    os.path.abspath("mous"),
    os.path.abspath(os.path.join(os.getcwd(), "mous")),
]

MOU_DIR: Optional[str] = None
for candidate in POSSIBLE_MOU_DIRS:
    if os.path.isdir(candidate):
        MOU_DIR = candidate
        break

# File mapping per role and optional subtype
DOCX_MAPPING: Dict[str, Dict[str, str]] = {
    "doctor": {
        "default": "CALLMEDEX_Doctors_Terms_and_Conditions_Modified.docx",
        "standard": "DOCTOR.docx",
    },
    "dentist": {
        "default": "CALLMEDEX_Dental_Clinics_Terms_and_Conditions.docx",
    },
    "dietitian": {
        "default": "CALLMEDEX_Dietetic_Services_Terms_and_Conditions.docx",
        "standard": "Dietician Terms & Conditions.docx",
    },
    "physiotherapist": {
        "default": "CALLMEDEX_Physiotherapy_Centers_and_Home_Services_Terms.docx",
        "standard": "PHYSIOTHERAPIST.docx",
    },
    "nurse": {
        "default": "CALLMEDEX_Nursing_Home_Services_Terms.docx",
        "standard": "Nursing.docx",
    },
    "phlebotomist": {
        "default": "FULL TIME PHLEBO (1).docx",
        "full_time": "FULL TIME PHLEBO (1).docx",
        "part_time": "PART TIME PHLEBO (1).docx",
    },
    "organization": {
        "default": "Diagnostic Center MOUs.docx",
        "diagnostic": "CALL MEDEX Diagnostic Services Agreement.docx",
        "ecg_xray": "ECG X RAY MOU.docx",
    },
    "processing_center": {
        "default": "Diagnostic Center MOUs.docx",
    },
    "staff": {
        "default": "LOGISTIC PERSON.docx",
        "logistics": "LOGISTIC PERSON.docx",
    },
    "pharmacy": {
        "default": "Diagnostic Center MOUs.docx",
    },
}

_CACHE: Dict[str, str] = {}


def _extract_docx_text(filepath: str) -> str:
    """Extract clean paragraphs and tables from a docx file."""
    try:
        import docx
        doc = docx.Document(filepath)
        sections = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                sections.append(text)
        for table in doc.tables:
            table_lines = []
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    table_lines.append(" | ".join(cells))
            if table_lines:
                sections.append("\n" + "\n".join(table_lines) + "\n")
        return "\n\n".join(sections)
    except Exception as e:
        logger.error(f"Failed to read docx {filepath}: {e}")
        return ""


def get_full_mou_text(role: str, subtype: Optional[str] = None) -> str:
    """
    Returns the complete, unabridged legal agreement for the specified role and subtype.
    """
    role_key = (role or "doctor").strip().lower()
    sub_key = (subtype or "default").strip().lower()

    cache_key = f"{role_key}:{sub_key}"
    if cache_key in _CACHE:
        return _CACHE[cache_key]

    mapping = DOCX_MAPPING.get(role_key, DOCX_MAPPING["doctor"])
    filename = mapping.get(sub_key) or mapping.get("default") or "DOCTOR.docx"

    if MOU_DIR:
        filepath = os.path.join(MOU_DIR, filename)
        if os.path.isfile(filepath):
            content = _extract_docx_text(filepath)
            if content and len(content) > 100:
                _CACHE[cache_key] = content
                return content

    # Fallback to standard if default not found
    if "standard" in mapping and MOU_DIR:
        alt_path = os.path.join(MOU_DIR, mapping["standard"])
        if os.path.isfile(alt_path):
            content = _extract_docx_text(alt_path)
            if content and len(content) > 100:
                _CACHE[cache_key] = content
                return content

    return ""
